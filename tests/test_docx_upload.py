"""
Comprehensive test suite for DOCX resume parsing, text extraction, validation,
error handling, and end-to-end route handling in the placement readiness agent.
"""

import io
import os
import pytest
from docx import Document
from utils.pdf_reader import (
    extract_text_from_docx,
    extract_text_from_file,
    ALLOWED_RESUME_EXTENSIONS,
    MAX_RESUME_TEXT_CHARS,
    MIN_RESUME_TEXT_CHARS,
    MIN_RESUME_WORD_COUNT,
)
from app import app


@pytest.fixture
def sample_docx_bytes():
    """Generates an in-memory .docx file with standard resume content."""
    doc = Document()
    doc.add_heading("Jane Doe - Software Engineer Resume", level=1)
    doc.add_paragraph("Email: jane.doe@example.com | Phone: +1 555-0199 | Location: San Francisco, CA")
    
    doc.add_heading("Summary", level=2)
    doc.add_paragraph(
        "Dedicated Computer Science graduate specializing in Python, Flask, Distributed Systems, "
        "and Machine Learning applications with comprehensive campus placement readiness."
    )
    
    doc.add_heading("Education", level=2)
    table = doc.add_table(rows=1, cols=3)
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = "Degree"
    hdr_cells[1].text = "Institution"
    hdr_cells[2].text = "CGPA"
    
    row_cells = table.add_row().cells
    row_cells[0].text = "B.Tech in Computer Engineering"
    row_cells[1].text = "National Institute of Technology"
    row_cells[2].text = "9.1 / 10.0"

    doc.add_heading("Technical Skills", level=2)
    doc.add_paragraph("Languages: Python, Java, C++, TypeScript, SQL")
    doc.add_paragraph("Frameworks & Tools: Flask, FastAPI, Docker, Git, PyTorch, PostgreSQL")

    doc.add_heading("Projects", level=2)
    doc.add_paragraph(
        "Autonomous Campus Placement Preparation Agent: Architected multi-agent LLM analysis "
        "pipeline for automated resume evaluation, benchmark scoring, and personalized study roadmaps."
    )

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf


@pytest.fixture
def empty_docx_bytes():
    """Generates an empty .docx document without paragraphs or content."""
    doc = Document()
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf


@pytest.fixture
def short_docx_bytes():
    """Generates a .docx file containing fewer words than MIN_RESUME_WORD_COUNT."""
    doc = Document()
    doc.add_paragraph("Short resume.")
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf


def test_allowed_resume_extensions_contains_docx_and_pdf():
    """Verify that both .pdf and .docx are in the allowed extensions set."""
    assert ".docx" in ALLOWED_RESUME_EXTENSIONS
    assert ".pdf" in ALLOWED_RESUME_EXTENSIONS


def test_extract_text_from_docx_success(sample_docx_bytes, tmp_path):
    """Verify successful extraction of headings, paragraphs, and tables from a DOCX file."""
    docx_file = tmp_path / "valid_resume.docx"
    docx_file.write_bytes(sample_docx_bytes.getvalue())

    extracted = extract_text_from_docx(str(docx_file))

    assert "Jane Doe - Software Engineer Resume" in extracted
    assert "Dedicated Computer Science graduate" in extracted
    assert "B.Tech in Computer Engineering" in extracted
    assert "National Institute of Technology" in extracted
    assert "Autonomous Campus Placement Preparation Agent" in extracted
    assert len(extracted) >= MIN_RESUME_TEXT_CHARS


def test_extract_text_from_docx_stream_support(sample_docx_bytes):
    """Verify that extract_text_from_docx accepts BytesIO buffer directly."""
    extracted = extract_text_from_docx(sample_docx_bytes)
    assert "Jane Doe" in extracted
    assert "Technical Skills" in extracted


def test_extract_text_from_docx_empty_file_fails(empty_docx_bytes, tmp_path):
    """Verify that extracting from an empty docx raises ValueError with clear feedback."""
    empty_file = tmp_path / "empty_resume.docx"
    empty_file.write_bytes(empty_docx_bytes.getvalue())

    with pytest.raises(ValueError) as excinfo:
        extract_text_from_docx(str(empty_file))
    assert "Could not extract sufficient text" in str(excinfo.value)


def test_extract_text_from_docx_short_content_fails(short_docx_bytes, tmp_path):
    """Verify that extracting insufficient resume text raises ValueError with guidance."""
    short_file = tmp_path / "short_resume.docx"
    short_file.write_bytes(short_docx_bytes.getvalue())

    with pytest.raises(ValueError) as excinfo:
        extract_text_from_docx(str(short_file))
    assert "Could not extract sufficient text" in str(excinfo.value)


def test_extract_text_from_docx_file_not_found():
    """Verify that passing a non-existent path raises FileNotFoundError."""
    missing_path = "non_existent_file_path_xyz_12345.docx"
    with pytest.raises(FileNotFoundError):
        extract_text_from_docx(missing_path)


def test_extract_text_from_docx_corrupt_file(tmp_path):
    """Verify that corrupt non-docx files raise ValueError for package errors."""
    corrupt_file = tmp_path / "corrupt_resume.docx"
    corrupt_file.write_bytes(b"This is just raw text, not a valid zip docx archive.")

    with pytest.raises(ValueError) as excinfo:
        extract_text_from_docx(str(corrupt_file))
    assert "not a valid or readable Word document" in str(excinfo.value)


def test_extract_text_from_docx_truncation(tmp_path):
    """Verify that excessively large DOCX documents are safely truncated to MAX_RESUME_TEXT_CHARS."""
    doc = Document()
    doc.add_heading("Massive Resume File", level=1)
    for i in range(150):
        doc.add_paragraph(
            f"Paragraph {i}: Engineering experience designing distributed systems, caching layers, "
            f"relational databases, message queues, and placement agent platforms for campus recruitment."
        )

    huge_file = tmp_path / "huge_resume.docx"
    doc.save(str(huge_file))

    extracted = extract_text_from_docx(str(huge_file))
    assert len(extracted) <= MAX_RESUME_TEXT_CHARS


def test_extract_text_from_file_routing(sample_docx_bytes, tmp_path):
    """Verify extract_text_from_file accurately routes .docx and handles unsupported types."""
    docx_file = tmp_path / "candidate_profile.docx"
    docx_file.write_bytes(sample_docx_bytes.getvalue())

    # Routes DOCX successfully
    extracted_docx = extract_text_from_file(str(docx_file))
    assert "Jane Doe" in extracted_docx

    # Unsupported extension rejection
    txt_file = tmp_path / "profile.txt"
    txt_file.write_text("Hello world testing unsupported format extension.")
    with pytest.raises(ValueError) as excinfo:
        extract_text_from_file(str(txt_file))
    assert "Unsupported file format" in str(excinfo.value)


def test_route_analyze_accepts_docx_upload(monkeypatch, sample_docx_bytes):
    """Verify the /analyze Flask endpoint processes .docx uploads and proceeds to dashboard."""
    # Mock coordinator agent to avoid real network LLM requests
    from agents.coordinator import CoordinatorAgent
    dummy_payload = {
        "resume_analysis": {"skills": ["Python", "Flask", "Docker"], "score": 85},
        "job_compatibility": {"match_percentage": 90, "matching_skills": ["Python", "Flask"]},
        "skill_gap": {"missing_skills": ["Kubernetes"], "recommended_skills": ["Go"]},
        "roadmap": {"weeks": [{"week": 1, "topic": "Advanced Python", "focus": "Core concepts"}]},
        "interview_preparation": {"questions": []},
    }
    monkeypatch.setattr(CoordinatorAgent, "analyze", lambda self, *args, **kwargs: dummy_payload)

    app.config["TESTING"] = True
    with app.test_client() as client:
        data = {
            "resume": (sample_docx_bytes, "candidate_cv.docx"),
            "company": "Google",
            "job_description": "We are looking for a Python backend engineer with campus placement readiness.",
        }
        res = client.post("/analyze", data=data, content_type="multipart/form-data", follow_redirects=False)
        assert res.status_code == 302
        assert "/dashboard/" in res.headers["Location"]


def test_route_analyze_rejects_unsupported_extensions():
    """Verify the /analyze Flask endpoint rejects non-pdf and non-docx files with an error flash."""
    app.config["TESTING"] = True
    with app.test_client() as client:
        fake_exe = io.BytesIO(b"executable dummy content")
        data = {
            "resume": (fake_exe, "malicious_script.exe"),
            "company": "General",
        }
        res = client.post("/analyze", data=data, content_type="multipart/form-data", follow_redirects=True)
        assert res.status_code == 200
        assert b"Only PDF (.pdf) and Word (.docx) resumes are supported" in res.data
